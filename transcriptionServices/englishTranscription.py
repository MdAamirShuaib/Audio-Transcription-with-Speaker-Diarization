import os
import shutil
import re
import requests
import pandas as pd
import docx
import random
import uuid
import configparser


def get_config():
    config = configparser.ConfigParser()
    config.read("config.ini", encoding="utf-8")
    return config


# --------------------Transcription Code for English Language---------------------
def read_file(filename, chunk_size=5242880):
    with open(filename, "rb") as _file:
        while True:
            data = _file.read(chunk_size)
            if not data:
                break
            yield data


def transcribe_english(file, token):
    headers = {"authorization": token}
    url_response = requests.post(
        "https://api.assemblyai.com/v2/upload", headers=headers, data=read_file(file)
    )

    url = url_response.json()["upload_url"]

    id_endpoint = "https://api.assemblyai.com/v2/transcript"

    json = {
        "audio_url": url,
        "speaker_labels": True,
        "auto_highlights": True,
        "iab_categories": True,
        "sentiment_analysis": True,
        "auto_chapters": True,
        "entity_detection": True,
    }

    headers = {"authorization": token, "content-type": "application/json"}
    id_response = requests.post(id_endpoint, json=json, headers=headers)
    transcribe_id = id_response.json()["id"]
    text_endpoint = "https://api.assemblyai.com/v2/transcript/" + transcribe_id
    headers = {
        "authorization": token,
    }
    result = requests.get(text_endpoint, headers=headers).json()

    while result.get("status") != "completed":
        if result.get("status") == "error":
            return "error"
        text_endpoint = f"https://api.assemblyai.com/v2/transcript/{transcribe_id}"
        headers = {"authorization": token}
        result = requests.get(text_endpoint, headers=headers).json()
    return result


# ------------------Extrcting Data from AssemblyAI's JSON response-------------------


def json_data_extraction(result, file_name):
    audindex = pd.json_normalize(result["words"])
    audindex["file_name"] = file_name
    chapters = pd.json_normalize(result["chapters"])
    topics = pd.json_normalize(result["iab_categories_result"]["results"])
    topics["label_1"] = topics["labels"].apply(lambda x: x[0]["label"])
    topics["label_2"] = topics["labels"].apply(
        lambda x: x[1]["label"] if len(x) > 1 else "none"
    )
    topics["label_3"] = topics["labels"].apply(
        lambda x: x[2]["label"] if len(x) > 2 else "none"
    )
    highlights = pd.json_normalize(result["auto_highlights_result"]["results"])
    highlights = highlights.text.unique()
    audindex["summary"] = ""
    audindex["headline"] = ""
    audindex["gist"] = ""
    audindex["label_1"] = ""
    audindex["label_2"] = ""
    audindex["label_3"] = ""
    speakers = list(audindex.speaker)
    previous_speaker = "A"
    l = len(speakers)
    i = 0
    speaker_seq_list = list()
    for index, new_speaker in enumerate(speakers):
        if index > 0:
            previous_speaker = speakers[index - 1]
        if new_speaker != previous_speaker:
            i += 1
        speaker_seq_list.append(i)
    audindex["sequence"] = speaker_seq_list

    for j in range(0, len(chapters)):
        for i in range(0, len(audindex)):
            if (
                audindex.iloc[i]["start"] >= chapters.iloc[j]["start"]
                and audindex.iloc[i]["end"] <= chapters.iloc[j]["end"]
            ):
                audindex.loc[i, "summary"] = chapters.iloc[j]["summary"]
                audindex.loc[i, "headline"] = chapters.iloc[j]["headline"]
                audindex.loc[i, "gist"] = chapters.iloc[j]["gist"]

    for j in range(0, len(topics)):
        try:
            for i in range(0, len(audindex)):
                if (
                    audindex.iloc[i]["start"] >= topics.iloc[j]["timestamp.start"]
                    and audindex.iloc[i]["end"] <= topics.iloc[j + 1]["timestamp.start"]
                ):
                    audindex.loc[i, "label_1"] = topics.iloc[j]["label_1"]
                    audindex.loc[i, "label_2"] = topics.iloc[j]["label_2"]
                    audindex.loc[i, "label_3"] = topics.iloc[j]["label_3"]
        except:
            for i in range(0, len(audindex)):
                if audindex.iloc[i]["start"] >= topics.iloc[j]["timestamp.start"]:
                    audindex.loc[i, "label_1"] = topics.iloc[j]["label_1"]
                    audindex.loc[i, "label_2"] = topics.iloc[j]["label_2"]
                    audindex.loc[i, "label_3"] = topics.iloc[j]["label_3"]

    group = [
        "file_name",
        "speaker",
        "summary",
        "headline",
        "gist",
        "label_1",
        "label_2",
        "label_3",
        "sequence",
    ]
    df = pd.DataFrame(
        audindex.groupby(group).agg(
            utter=("text", " ".join),
            start_time=("start", "min"),
            end_time=("end", "max"),
        )
    )
    df.reset_index(inplace=True)
    df["key_phrase"] = "none"
    for x in highlights:
        df.loc[(df.utter.str.contains(x)), "key_phrase"] = x
    df.sort_values(by=["start_time"], inplace=True)

    group_doc = ["speaker", "sequence"]
    df_doc = pd.DataFrame(
        audindex.groupby(group_doc).agg(
            utter=("text", " ".join), start_time=("start", "min")
        )
    )
    df_doc.reset_index(inplace=True)
    df_doc.sort_values(by=["start_time"], inplace=True)
    df_doc_grouped = df_doc[["speaker", "utter"]]

    doc = docx.Document()
    t = doc.add_table(df_doc_grouped.shape[0] + 1, df_doc_grouped.shape[1])
    for j in range(df_doc_grouped.shape[-1]):
        t.cell(0, j).text = df_doc_grouped.columns[j]
    for i in range(df_doc_grouped.shape[0]):
        for j in range(df_doc_grouped.shape[-1]):
            t.cell(i + 1, j).text = str(df_doc_grouped.values[i, j])

    return df, doc


async def start_transcription(files):
    config = get_config()
    tokens = eval(config["AssemblyAI"]["tokens"])
    os.makedirs("documents", exist_ok=True)
    data = []
    errors = []
    try:
        for file in files:
            rawfile_name = file.filename
            file_name = re.sub("[^a-zA-Z0-9.@()!_-]", "", rawfile_name)[:-4]
            os.makedirs("documents/audiofile", exist_ok=True)
            filepath = "documents/audiofile/" + file.filename
            with open(filepath, "wb") as audio_file_object:
                shutil.copyfileobj(file.file, audio_file_object)
            audio_file_object.close()
            token = tokens[random.randint(0, len(tokens) - 1)]
            result = transcribe_english(filepath, token)
            if result == "error":
                errors.append("Transcription Error" + file_name)
                os.remove(filepath)
                continue
            print("transcription completed")
            df, doc = json_data_extraction(result, file_name)
            data.append(df.to_json(orient="records"))
            os.remove(filepath)

    except Exception as e:
        print(e)
        errors.append("Unexpected error occurred during transcribing")
    return errors, data
